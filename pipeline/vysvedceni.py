"""Vysvědčení koalici 2022–2026 — převod ručního auditu do dat webu.

TENHLE MODUL NIC NESKLÍZÍ. Zdrojem je sešit `podklady/vysvedceni-koalice-2022-2026.xlsx`,
ruční audit: člověk prošel všech 38 projektů z portálu muml.pincity.cz/projekty
a porovnal je s registrem smluv, věstníkem zakázek, registry dotací, závěrečnými
účty a médii. Modul ho jen převádí do JSON, aby web stavěl ze stejných dat jako
zbytek přehledu a aby šlo ověřit, že se cestou nic nezměnilo.

PROČ SE TO ODDĚLUJE OD SLIBNÍKU. `/slibnik` má tvrdou zásadu: nikdy netvrdit
„nesplněno", protože u úkolů z usnesení je výsledek kontroly vždy jen v příloze,
kterou město nezveřejňuje — nevíme to a nemůžeme to vědět. Tady je situace jiná:
tvrzení se ověřuje proti NEZÁVISLÝM REGISTRŮM. „V registru smluv není k tomuhle
projektu žádná smlouva" je zjištění o registru, ne dohad o příloze.

Ta zásada ale platí dál tam, kde se opírá o nenález obecně. Proto se u slibů
z programového prohlášení rozlišuje:

  * `nesplneno`  — doloženo pozitivně (rozpočtová položka čerpána na 0 Kč apod.)
  * `nedohledano` — nepodařilo se najít veřejný doklad o plnění

Obojí má v sešitu známku 5, ale NENÍ to totéž a web to nesmí slít dohromady:
sedm z dvanácti „pětek" je nedohledáno. Kdyby se to sečetlo, web by tvrdil víc,
než audit unese — a přesně to si tenhle projekt zakazuje.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from lib.core import Log, uloz  # noqa: E402

VYSTUP = "vysvedceni/audit.json"

# Sešit leží rovnou v `web/public/`, odkud ho web servíruje ke stažení.
# Druhá kopie mimo web by se dřív nebo později rozešla s tou publikovanou
# a čtenář by si pak ověřoval jiná čísla, než jaká na stránce vidí.
SESIT_VEREJNE = "/podklady/vysvedceni-koalice-2022-2026.xlsx"
SESIT = Path(__file__).resolve().parent.parent / "web" / "public" / SESIT_VEREJNE.lstrip("/")

ZNAMKY = {
    1: "Splněno",
    2: "Splněno s výhradami",
    3: "Skluz nebo neúplné plnění",
    4: "Jen na papíře nebo zavádějící deklarace",
    5: "Nesplněno",
}

# Stavy slibů z programového prohlášení roztříděné na to, co je DOLOŽENO,
# a co se jen nepodařilo dohledat. Klíč je doslovný text ze sešitu.
NEDOHLEDANO = {"Nedohledáno", "Nedohledáno / nesplněno"}


def _text(c) -> str | None:
    if c is None:
        return None
    s = str(c).strip()
    return s or None


def _cislo(c) -> float | None:
    if c is None or isinstance(c, str) and not c.strip():
        return None
    try:
        return float(c)
    except (TypeError, ValueError):
        return None


def _zdroje(s: str | None) -> list[str]:
    """Rozdělí buňku se zdroji na jednotlivé položky.

    Zdroje jsou v sešitu oddělené středníkem nebo prostředníkem. Nechávají se
    jako text — je mezi nimi směs URL, spisových značek a názvů dokumentů
    a vyrábět z nenapsaného odkaz by znamenalo si ho vymyslet.
    """
    if not s:
        return []
    kusy = re.split(r"[;·]\s*", s)
    return [k.strip() for k in kusy if k.strip()]


def _projekty(ws) -> list[dict]:
    """Řádky projektů z listu Vysvědčení.

    Pod tabulkou je souhrnný řádek CELKEM a za ním známkovací stupnice, jejíž
    řádky mají v prvním sloupci taky číslo (1 až 5). Kdyby se četlo jen „první
    sloupec je číslo", stupnice by se do dat dostala jako pět dalších projektů
    a souhrn by hlásil 43 místo 38. Proto se čte jen po CELKEM.
    """
    out = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        nazev = _text(r[1])
        if nazev and (nazev.startswith("CELKEM") or nazev == "Známkovací stupnice"):
            break
        ident = _text(r[0])
        if not ident or not ident.isdigit():
            continue
        znamka = _cislo(r[11])
        out.append({
            "id": int(ident),
            "nazev": _text(r[1]),
            "tema": _text(r[2]),
            "typ": _text(r[3]),
            "deklarovany_stav": _text(r[4]),
            "deklarovany_termin": _text(r[5]),
            "deklarovany_rozpocet_czk": _cislo(r[6]),
            "deklarovana_dotace_czk": _cislo(r[7]),
            "skutecny_stav": _text(r[8]),
            "dolozene_naklady_czk": _cislo(r[9]),
            "dolozena_dotace_czk": _cislo(r[10]),
            "znamka": int(znamka) if znamka else None,
            "hodnoceni": _text(r[12]),
            "rozpor": _text(r[13]),
            "jistota": _text(r[14]),
            "zdroje": _zdroje(_text(r[15])),
        })
    return out


def _sliby(ws) -> list[dict]:
    out = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        oblast, slib, stav = _text(r[0]), _text(r[1]), _text(r[2])
        if not slib or stav == "Stav":
            continue
        znamka = _cislo(r[4])
        out.append({
            "oblast": oblast,
            "slib": slib,
            "stav": stav,
            # Rozlišení, na kterém stojí poctivost celé sekce.
            "dolozeno": stav not in NEDOHLEDANO,
            "doklad": _text(r[3]),
            "znamka": int(znamka) if znamka else None,
        })
    return out


def _mimo_portal(ws) -> list[dict]:
    out = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        tema = _text(r[0])
        if not tema or tema == "Téma":
            continue
        out.append({
            "tema": tema,
            "slibeno": _text(r[1]),
            "stalo_se": _text(r[2]),
            "hodnoceni": _text(r[3]),
            "zdroje": _zdroje(_text(r[4])),
        })
    return out


def _rozpocet(ws) -> dict:
    """Kapitálové výdaje po letech a akce s nulovým čerpáním.

    Poslední řádek tabulky let je souhrn za 2023–2025, ne rok. Kdyby zůstal
    mezi lety, graf by ho nakreslil jako šestý sloupec, který je součtem tří
    ostatních — a čtenář by viděl rok, který neexistuje. Proto stojí zvlášť.
    """
    roky, nulove = [], []
    souhrn = None
    rezim = "roky"
    for r in ws.iter_rows(min_row=3, values_only=True):
        prvni = _text(r[0])
        if not prvni:
            continue
        if prvni.startswith("Investiční akce"):
            rezim = "nulove"
            continue
        if prvni == "Rok" or prvni.startswith("Schválený rozpočet 2026"):
            continue
        if rezim == "roky":
            skutecnost = _cislo(r[3])
            if skutecnost is None:
                continue
            zaznam = {
                "rok": prvni,
                "schvaleno_tis": _cislo(r[1]),
                "upraveno_tis": _cislo(r[2]),
                "skutecnost_tis": skutecnost,
                "plneni_upraveny": _cislo(r[4]),
                "nevycerpano_tis": _cislo(r[6]),
            }
            if prvni.lower().startswith(("celkem", "souhrn")):
                souhrn = zaznam
            else:
                roky.append(zaznam)
        else:
            akce = _text(r[1])
            if not akce:
                continue
            nulove.append({
                "rok": prvni,
                "akce": akce,
                "upraveno_tis": _cislo(r[2]),
                "skutecnost_tis": _cislo(r[3]),
                "cerpani": _cislo(r[4]),
            })
    return {"roky": roky, "souhrn_obdobi": souhrn, "nulove_cerpani": nulove}


def _metodika(ws) -> list[dict]:
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        nadpis, text = _text(r[0]), _text(r[1])
        if nadpis and text:
            out.append({"nadpis": nadpis, "text": text})
    return out


def main() -> dict:
    log = Log("vysvedceni")
    try:
        import openpyxl
    except ImportError:
        log.chyba("openpyxl není nainstalované — `pip install openpyxl`")
        return log.uzavri()

    if not SESIT.exists():
        log.chyba(f"podkladový sešit chybí: {SESIT}")
        return log.uzavri()

    wb = openpyxl.load_workbook(SESIT, data_only=True)
    projekty = _projekty(wb["Vysvědčení"])
    sliby = _sliby(wb["Programové prohlášení"])

    if not projekty:
        log.chyba("v sešitu nejsou žádné projekty — změnila se struktura listu?")
        return log.uzavri()

    # Souhrn se POČÍTÁ tady, nepřepisuje se z listu Souhrn. Kdyby se čísla
    # v sešitu rozešla s řádky, web má ukazovat to, co v řádcích doopravdy je.
    podle_znamky = {z: sum(1 for p in projekty if p["znamka"] == z) for z in ZNAMKY}
    dolozene_sliby = [s for s in sliby if s["dolozeno"]]
    nedohledane = [s for s in sliby if not s["dolozeno"]]

    vystup = {
        "hodnoceno_k": "2026-08-27",
        "volebni_obdobi": "2022–2026",
        "predmet": "38 projektů zveřejněných na portálu muml.pincity.cz/projekty",
        "sesit": SESIT_VEREJNE,
        "znamky": ZNAMKY,
        "souhrn": {
            "projektu": len(projekty),
            "podle_znamky": podle_znamky,
            "deklarovany_rozpocet_czk": sum(p["deklarovany_rozpocet_czk"] or 0 for p in projekty),
            "deklarovana_dotace_czk": sum(p["deklarovana_dotace_czk"] or 0 for p in projekty),
            "dolozene_naklady_czk": sum(p["dolozene_naklady_czk"] or 0 for p in projekty),
            "dolozena_dotace_czk": sum(p["dolozena_dotace_czk"] or 0 for p in projekty),
            "slibu": len(sliby),
            "slibu_dolozenych": len(dolozene_sliby),
            "slibu_nedohledanych": len(nedohledane),
        },
        "projekty": projekty,
        "sliby": sliby,
        "mimo_portal": _mimo_portal(wb["Témata mimo portál"]),
        "rozpocet": _rozpocet(wb["Rozpočet a investice"]),
        "metodika": _metodika(wb["Metodika a zdroje"]),
    }

    uloz(VYSTUP, vystup)
    log.info("projekty", pocet=len(projekty),
             znamky=" ".join(f"{z}:{n}" for z, n in podle_znamky.items()))
    log.info("sliby z programového prohlášení",
             celkem=len(sliby), dolozeno=len(dolozene_sliby), nedohledano=len(nedohledane))
    log.info("uloženo", soubor=VYSTUP)
    log.pricti(len(projekty))
    return log.uzavri()


if __name__ == "__main__":
    main()
